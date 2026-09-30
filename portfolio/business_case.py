import pandas as pd
import numpy as np
from config import BUSINESS_CASE


def compute_business_case_a(df_rfm: pd.DataFrame, config: dict = BUSINESS_CASE) -> dict:
    df_at_risk = df_rfm[df_rfm["segment"] == "À risque"].copy()
    
    nb_client = df_at_risk["CustomerID"].nunique()
    n_control = int(round(nb_client * config["control_group_pct"]))
    n_treatment = nb_client - n_control

    cost = float(round(n_treatment * config["contact_cost_gbp"], 2))
    valeur_commande = float(round((df_at_risk["Monetary"] / df_at_risk["Frequency"]).mean(), 2))
    
    n_responder = n_treatment * config["response_rate"]
    gross_margin = float(round(n_responder * valeur_commande * config["margin_rate"], 2))
    net_margin = float(round(gross_margin - cost, 2))

    break_event_rate = float(round(config["contact_cost_gbp"] / (valeur_commande * config["margin_rate"]), 4))
    roi = float(round((net_margin / cost) * 100, 1)) if cost > 0 else 0.0

    return {
        "nb_client": nb_client,
        "n_control": n_control,
        "n_treatment": n_treatment,
        "cost": cost,
        "valeur_commande": valeur_commande,
        "n_responder": round(n_responder, 1),
        "gross_margin": gross_margin,
        "net_margin": net_margin,
        "break_event_rate": break_event_rate,
        "roi_pct": roi,
    }


def compute_business_case_b(df_candidates: pd.DataFrame, config: dict = BUSINESS_CASE) -> dict:
    nb_candidates = len(df_candidates)
    ca_at_risk = float(round(df_candidates["ca_total"].sum(), 2))
    
    lost_margin = float(round(ca_at_risk * (1 - config["transfer_rate"]) * config["margin_rate"], 2))
    saving = float(round(nb_candidates * config["cost_per_sku_gbp"], 2))
    net_gain = float(round(saving - lost_margin, 2))
    break_even_sku_cost = float(round(lost_margin / nb_candidates, 2)) if nb_candidates > 0 else 0.0

    return {
        "nb_candidates": nb_candidates,
        "ca_at_risk": ca_at_risk,
        "lost_margin": lost_margin,
        "saving": saving,
        "net_gain": net_gain,
        "break_even_sku_cost": break_even_sku_cost,
    }


def compute_sensitivity_tables(
    df_rfm: pd.DataFrame, 
    df_candidates: pd.DataFrame, 
    config: dict = BUSINESS_CASE
) -> tuple[pd.DataFrame, pd.DataFrame]:
    df_at_risk = df_rfm[df_rfm["segment"] == "À risque"]
    n_treatment = len(df_at_risk) * (1 - config["control_group_pct"])
    valeur_commande = (df_at_risk["Monetary"] / df_at_risk["Frequency"]).mean()
    cost_a = n_treatment * config["contact_cost_gbp"]

    # Table Sensibilité Case A : Response Rate vs Margin Rate
    response_rates = [0.02, 0.04, 0.06, 0.08, 0.10, 0.12, 0.15]
    margin_rates = [0.20, 0.25, 0.30, 0.35, 0.40, 0.50]
    
    sens_a = pd.DataFrame(index=[f"{r*100:.0f}%" for r in response_rates], columns=[f"{m*100:.0f}%" for m in margin_rates])
    for r in response_rates:
        for m in margin_rates:
            gain = (n_treatment * r * valeur_commande * m) - cost_a
            sens_a.loc[f"{r*100:.0f}%", f"{m*100:.0f}%"] = float(round(gain, 2))

    # Table Sensibilité Case B : Transfer Rate vs Cost per SKU
    transfer_rates = [0.10, 0.30, 0.50, 0.70, 0.90]
    sku_costs = [100, 300, 500, 750, 1000]
    ca_at_risk = df_candidates["ca_total"].sum()
    nb_candidates = len(df_candidates)
    
    sens_b = pd.DataFrame(index=[f"{t*100:.0f}%" for t in transfer_rates], columns=[f"{c} £" for c in sku_costs])
    for t in transfer_rates:
        for c in sku_costs:
            gain = (nb_candidates * c) - (ca_at_risk * (1 - t) * config["margin_rate"])
            sens_b.loc[f"{t*100:.0f}%", f"{c} £"] = float(round(gain, 2))

    return sens_a, sens_b


def generate_business_case_excel(df_rfm: pd.DataFrame, df_candidates: pd.DataFrame, config: dict = BUSINESS_CASE, output_path: str = "outputs/portfolio/business_case.xlsx") -> str:
    """Génère un fichier Excel dynamique avec des valeurs initiales calculées et des formules natives."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    
    # Style definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
    bold_font = Font(name="Calibri", size=11, bold=True)
    num_fmt_curr = "£#,##0.00"
    num_fmt_pct = "0.0%"
    num_fmt_int = "#,##0"

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # -------------------------------------------------------------
    # Onglet 1 : Hypothèses (Assumptions)
    # -------------------------------------------------------------
    ws_hyp = wb.active
    ws_hyp.title = "Hypothèses"
    ws_hyp.views.sheetView[0].showGridLines = True

    ws_hyp["A1"] = "CUSTOMER & GROWTH AI LAB - HYPOTHÈSES DES BUSINESS CASES"
    ws_hyp["A1"].font = title_font
    ws_hyp["A2"] = "Modifiez les cellules en fond vert léger ci-dessous pour recalculer dynamiquement les modèles."
    ws_hyp["A2"].font = subtitle_font

    headers_hyp = ["Code Hypothèse", "Description de l'Hypothèse", "Valeur", "Unité", "Source / Validation"]
    for col_num, h in enumerate(headers_hyp, 1):
        cell = ws_hyp.cell(row=4, column=col_num)
        cell.value = h
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    hyp_data = [
        ("COST_CONTACT", "Coût unitaire de contact CRM (Phoning / SMS)", 2.00, "£ / client", "Estimation CRM direct"),
        ("RESP_RATE", "Taux de conversion / réengagement estimé", 0.08, "%", "Hypothèse de travail à valider par A/B testing"),
        ("MARGIN_RATE", "Taux de marge brute moyen sur les ventes", 0.35, "%", "Comptabilité analytique"),
        ("SKU_COST", "Coût de complexité logistique annuel par SKU", 500.00, "£ / SKU / an", "Analyse coûts fixes Supply"),
        ("TRANSFER_RATE", "Taux de transfert d'achat vers SKUs A/B", 0.50, "%", "Hypothèse de travail à valider par A/B testing"),
        ("CONTROL_PCT", "Part des clients réservés au groupe de contrôle", 0.10, "%", "Méthodologie AB Testing")
    ]

    input_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    for idx, (code, desc, val, unit, src) in enumerate(hyp_data, start=5):
        ws_hyp.cell(row=idx, column=1, value=code).font = Font(name="Calibri", bold=True)
        ws_hyp.cell(row=idx, column=2, value=desc)
        
        val_cell = ws_hyp.cell(row=idx, column=3, value=val)
        val_cell.fill = input_fill
        val_cell.font = Font(name="Calibri", bold=True)
        if unit == "£" or "£" in unit:
            val_cell.number_format = num_fmt_curr
        elif unit == "%":
            val_cell.number_format = num_fmt_pct

        ws_hyp.cell(row=idx, column=4, value=unit).alignment = Alignment(horizontal="center")
        ws_hyp.cell(row=idx, column=5, value=src)

        for col in range(1, 6):
            ws_hyp.cell(row=idx, column=col).border = thin_border

    # -------------------------------------------------------------
    # Onglet 2 : Business Case A (Valeurs & Formules Excel)
    # -------------------------------------------------------------
    ws_a = wb.create_sheet(title="Business Case A (Reconquête)")
    ws_a.views.sheetView[0].showGridLines = True

    df_at_risk = df_rfm[df_rfm["segment"] == "À risque"]
    nb_at_risk = df_at_risk["CustomerID"].nunique()
    valeur_commande = float(round((df_at_risk["Monetary"] / df_at_risk["Frequency"]).mean(), 2))

    ws_a["A1"] = "BUSINESS CASE A : RECONQUÊTE DES CLIENTS À RISQUE"
    ws_a["A1"].font = title_font
    ws_a["A2"] = "La colonne B affiche les valeurs initiales calculées. La colonne C contient les formules dynamiques liées à l'onglet 'Hypothèses'."
    ws_a["A2"].font = subtitle_font

    headers_a = ["Indicateur Financier / Opérationnel", "Valeur Initiale Calculée", "Formule Excel Dynamique", "Unité"]
    for col_num, h in enumerate(headers_a, 1):
        cell = ws_a.cell(row=4, column=col_num)
        cell.value = h
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    n_ctrl = int(round(nb_at_risk * config["control_group_pct"]))
    n_treat = nb_at_risk - n_ctrl
    c_contact = config["contact_cost_gbp"]
    c_total = n_treat * c_contact
    r_rate = config["response_rate"]
    n_resp = n_treat * r_rate
    m_rate = config["margin_rate"]
    ca_gross = n_resp * valeur_commande
    margin_gross = ca_gross * m_rate
    net_margin_a = margin_gross - c_total
    be_rate_a = c_contact / (valeur_commande * m_rate)
    roi_a = (net_margin_a / c_total) if c_total > 0 else 0.0

    rows_a = [
        ("Volume de clients ciblés (À risque)", nb_at_risk, f"={nb_at_risk}", num_fmt_int, "clients"),
        ("Part du groupe de contrôle", config["control_group_pct"], "='Hypothèses'!C10", num_fmt_pct, "%"),
        ("Volume du groupe de contrôle", n_ctrl, "=ROUND(C5*C6, 0)", num_fmt_int, "clients (AB Testing)"),
        ("Volume réellement contacté (Traitement)", n_treat, "=C5-C7", num_fmt_int, "clients"),
        ("Coût unitaire de contact", c_contact, "='Hypothèses'!C5", num_fmt_curr, "£ / client"),
        ("Investissement Total (Coût Campagne)", c_total, "=C8*C9", num_fmt_curr, "£"),
        ("Valeur Moyenne d'une Commande", valeur_commande, f"={valeur_commande}", num_fmt_curr, "£ / commande"),
        ("Taux de Réponse (Conversion)", r_rate, "='Hypothèses'!C6", num_fmt_pct, "%"),
        ("Nombre de clients réengagés", n_resp, "=C8*C12", num_fmt_int, "clients"),
        ("Chiffre d'Affaires Brut Généré", ca_gross, "=C13*C11", num_fmt_curr, "£"),
        ("Taux de Marge Brute", m_rate, "='Hypothèses'!C7", num_fmt_pct, "%"),
        ("Marge Brute Générée", margin_gross, "=C14*C15", num_fmt_curr, "£"),
        ("GAIN NET FINANCIER (Marge - Coût)", net_margin_a, "=C16-C10", num_fmt_curr, "£"),
        ("Seuil de Rentabilité (Break-even Rate)", be_rate_a, "=C9/(C11*C15)", num_fmt_pct, "% réengagement minimum"),
        ("RETOUR SUR INVESTISSEMENT (ROI)", roi_a, "=C17/C10", num_fmt_pct, "% ROI Net")
    ]

    for idx, (label, val_calc, formula_val, fmt, unit) in enumerate(rows_a, start=5):
        cell_lbl = ws_a.cell(row=idx, column=1, value=label)
        cell_val = ws_a.cell(row=idx, column=2, value=val_calc)
        cell_form = ws_a.cell(row=idx, column=3, value=formula_val)
        cell_unit = ws_a.cell(row=idx, column=4, value=unit)
        
        cell_val.number_format = fmt
        cell_form.number_format = fmt

        if "GAIN NET" in label or "RETOUR SUR INVESTISSEMENT" in label:
            cell_lbl.font = bold_font
            cell_val.font = bold_font
            cell_val.fill = input_fill
            cell_form.font = bold_font

        for c in (cell_lbl, cell_val, cell_form, cell_unit):
            c.border = thin_border

    # -------------------------------------------------------------
    # Onglet 3 : Business Case B (Valeurs & Formules Excel)
    # -------------------------------------------------------------
    ws_b = wb.create_sheet(title="Business Case B (SKUs)")
    ws_b.views.sheetView[0].showGridLines = True

    nb_sku_candidates = len(df_candidates)
    ca_candidates = float(round(df_candidates["ca_total"].sum(), 2))
    t_rate = config["transfer_rate"]
    lost_pct = 1.0 - t_rate
    ca_lost = ca_candidates * lost_pct
    margin_lost = ca_lost * m_rate
    sku_cost_val = config["cost_per_sku_gbp"]
    saving_b = nb_sku_candidates * sku_cost_val
    net_gain_b = saving_b - margin_lost
    be_sku_cost = margin_lost / nb_sku_candidates if nb_sku_candidates > 0 else 0.0

    ws_b["A1"] = "BUSINESS CASE B : RATIONALISATION DU CATALOGUE SKUS"
    ws_b["A1"].font = title_font
    ws_b["A2"] = "La colonne B affiche les valeurs initiales calculées. La colonne C contient les formules dynamiques liées à l'onglet 'Hypothèses'."
    ws_b["A2"].font = subtitle_font

    headers_b = ["Indicateur Financier / Logistique", "Valeur Initiale Calculée", "Formule Excel Dynamique", "Unité"]
    for col_num, h in enumerate(headers_b, 1):
        cell = ws_b.cell(row=4, column=col_num)
        cell.value = h
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rows_b = [
        ("Volume de SKUs C déréférencés", nb_sku_candidates, f"={nb_sku_candidates}", num_fmt_int, "références C"),
        ("Chiffre d'Affaires Historique des SKUs", ca_candidates, f"={ca_candidates}", num_fmt_curr, "£"),
        ("Taux de transfert d'achat vers SKUs A/B", t_rate, "='Hypothèses'!C9", num_fmt_pct, "%"),
        ("Part de CA perdu définitivement", lost_pct, "=1-C7", num_fmt_pct, "%"),
        ("Chiffre d'Affaires Perdu", ca_lost, "=C6*C8", num_fmt_curr, "£"),
        ("Taux de Marge Brute", m_rate, "='Hypothèses'!C7", num_fmt_pct, "%"),
        ("Marge Brute Perdue", margin_lost, "=C9*C10", num_fmt_curr, "£ (Coût d'opportunité)"),
        ("Coût annuel de complexité par SKU", sku_cost_val, "='Hypothèses'!C8", num_fmt_curr, "£ / SKU / an"),
        ("Économies Logistiques Annuelles", saving_b, "=C5*C12", num_fmt_curr, "£ de coûts fixes économisés"),
        ("GAIN NET FINANCIER ANNUEL", net_gain_b, "=C13-C11", num_fmt_curr, "£ Gain Net"),
        ("Seuil de Rentabilité (Coût Logistique min / SKU)", be_sku_cost, "=C11/C5", num_fmt_curr, "£ / SKU / an min")
    ]

    for idx, (label, val_calc, formula_val, fmt, unit) in enumerate(rows_b, start=5):
        cell_lbl = ws_b.cell(row=idx, column=1, value=label)
        cell_val = ws_b.cell(row=idx, column=2, value=val_calc)
        cell_form = ws_b.cell(row=idx, column=3, value=formula_val)
        cell_unit = ws_b.cell(row=idx, column=4, value=unit)
        
        cell_val.number_format = fmt
        cell_form.number_format = fmt

        if "GAIN NET" in label or "Économies Logistiques" in label:
            cell_lbl.font = bold_font
            cell_val.font = bold_font
            cell_val.fill = input_fill
            cell_form.font = bold_font

        for c in (cell_lbl, cell_val, cell_form, cell_unit):
            c.border = thin_border


    # -------------------------------------------------------------
    # Onglet 4 : Liste détaillée des Clients À Risque (Data)
    # -------------------------------------------------------------
    ws_clients = wb.create_sheet(title="Clients À Risque (Data)")
    ws_clients.views.sheetView[0].showGridLines = True

    ws_clients["A1"] = "LISTE NATIVE DES CLIENTS DU SEGMENT À RISQUE (828 CLIENTS)"
    ws_clients["A1"].font = title_font

    headers_clients = ["CustomerID", "Récence (Jours)", "Fréquence (Achats)", "Montant Total (£)", "Segment RFM"]
    for col_num, h in enumerate(headers_clients, 1):
        cell = ws_clients.cell(row=3, column=col_num)
        cell.value = h
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, r in enumerate(df_at_risk.itertuples(), start=4):
        ws_clients.cell(row=idx, column=1, value=str(getattr(r, "CustomerID", "")))
        ws_clients.cell(row=idx, column=2, value=int(getattr(r, "Recency", 0))).number_format = num_fmt_int
        ws_clients.cell(row=idx, column=3, value=int(getattr(r, "Frequency", 0))).number_format = num_fmt_int
        ws_clients.cell(row=idx, column=4, value=float(getattr(r, "Monetary", 0.0))).number_format = num_fmt_curr
        ws_clients.cell(row=idx, column=5, value=str(getattr(r, "segment", "À risque")))

    # -------------------------------------------------------------
    # Onglet 5 : Liste détaillée des SKUs Candidats (Data)
    # -------------------------------------------------------------
    ws_skus = wb.create_sheet(title="SKUs Candidats (Data)")
    ws_skus.views.sheetView[0].showGridLines = True

    ws_skus["A1"] = "LISTE NATIVE DES 1 737 SKUS C CANDIDATS À LA DÉRÉFÉRENCIATION"
    ws_skus["A1"].font = title_font

    headers_skus = ["StockCode", "Description", "Chiffre d'Affaires (£)", "Quantité Vendue", "Classe Pareto ABC"]
    for col_num, h in enumerate(headers_skus, 1):
        cell = ws_skus.cell(row=3, column=col_num)
        cell.value = h
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, r in enumerate(df_candidates.itertuples(), start=4):
        ws_skus.cell(row=idx, column=1, value=str(getattr(r, "StockCode", "")))
        ws_skus.cell(row=idx, column=2, value=str(getattr(r, "Description", "N/A")))
        ws_skus.cell(row=idx, column=3, value=float(getattr(r, "ca_total", 0.0))).number_format = num_fmt_curr
        ws_skus.cell(row=idx, column=4, value=int(getattr(r, "quantity_total", 0))).number_format = num_fmt_int
        ws_skus.cell(row=idx, column=5, value=str(getattr(r, "categorie_abc", "C")))

    # Ajustement automatique des largeurs de colonnes
    for ws in [ws_hyp, ws_a, ws_b, ws_clients, ws_skus]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col[:50]) # échantillon pour performance
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    from pathlib import Path
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path



